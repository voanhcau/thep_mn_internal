from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "document"
ASSET_DIR = ROOT / ".docx-assets"
OUTPUT = OUT_DIR / "Huong_dan_kien_truc_va_xac_thuc_API_han_muc_tin_dung.docx"

NAVY = "16324F"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
DARK = "202124"
WHITE = "FFFFFF"
GREEN = "1F6B4F"
PALE_GREEN = "EAF5EF"
GOLD = "7A5A00"
PALE_GOLD = "FFF7E0"
RED = "9B1C1C"
PALE_RED = "FDECEC"
MONO = "Consolas"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Calibri", size=11, color=DARK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_number_instance(doc, num_id, abstract_id):
    numbering = doc.part.numbering_part.element
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    apply_num(p, 1)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        set_run_font(p.add_run(text))
    return p


def add_step(doc, title, detail, num_id=10):
    p = doc.add_paragraph()
    apply_num(p, num_id)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    set_run_font(p.add_run(title + ": "), bold=True, color=NAVY)
    set_run_font(p.add_run(detail))
    return p


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.10)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F6F8FA")
    p_pr.append(shd)
    run = p.add_run(code)
    set_run_font(run, name=MONO, size=8.5, color="24292F")
    return p


def add_callout(doc, label, text, kind="info"):
    fills = {"info": PALE_BLUE, "ok": PALE_GREEN, "warn": PALE_GOLD, "risk": PALE_RED}
    colors = {"info": NAVY, "ok": GREEN, "warn": GOLD, "risk": RED}
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fills[kind])
    p_pr.append(shd)
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(label + "  "), bold=True, color=colors[kind])
    set_run_font(p.add_run(text), color=DARK)


def add_table(doc, headers, rows, widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_cant_split(hdr)
    for idx, value in enumerate(headers):
        set_cell_shading(hdr.cells[idx], LIGHT_BLUE)
        p = hdr.cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size=font_size, bold=True, color=NAVY)
    for row_values in rows:
        row = table.add_row()
        set_cant_split(row)
        for idx, value in enumerate(row_values):
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            set_run_font(p.add_run(str(value)), size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_para(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        set_run_font(p.add_run(text), italic=italic)
    return p


def create_layer_diagram(path):
    img = Image.new("RGB", (1500, 910), "white")
    draw = ImageDraw.Draw(img)
    try:
        regular = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 31)
        bold = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 34)
        small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 25)
    except OSError:
        regular = bold = small = ImageFont.load_default()
    layers = [
        ("API", "HTTP, model binding, authorization, response", "DCEAF7"),
        ("APPLICATION", "Use case, validation, repository interface", "E8EEF5"),
        ("DOMAIN", "CreditLimit - business data model", "EDF5F1"),
        ("INFRASTRUCTURE", "Dapper, SqlConnection, parameterized SQL", "FFF4D6"),
        ("LEGACY SQL SERVER", "dbo.vw_Tin_Dung", "F4EAEA"),
    ]
    y = 55
    for i, (name, detail, fill) in enumerate(layers):
        draw.rounded_rectangle((120, y, 1380, y + 125), radius=18, fill="#" + fill, outline="#2E74B5", width=3)
        draw.text((165, y + 22), name, font=bold, fill="#16324F")
        draw.text((690, y + 27), detail, font=regular, fill="#202124")
        if i < len(layers) - 1:
            draw.line((750, y + 126, 750, y + 158), fill="#667085", width=5)
            draw.polygon([(738, y + 148), (762, y + 148), (750, y + 165)], fill="#667085")
        y += 165
    draw.text((120, 875), "Dependency direction: API -> Application <- Infrastructure; Domain contains no database code.", font=small, fill="#667085")
    img.save(path)


def create_auth_diagram(path):
    img = Image.new("RGB", (1600, 920), "white")
    draw = ImageDraw.Draw(img)
    try:
        regular = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 28)
        bold = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 32)
        small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 24)
    except OSError:
        regular = bold = small = ImageFont.load_default()

    boxes = {
        "client": (80, 320, 470, 590, "CLIENT / ODOO", "1. Xin token\n4. Gọi API với Bearer token", "E8EEF5"),
        "idp": (605, 85, 1030, 360, "AUTHORIZATION SERVER", "Giữ PRIVATE KEY\nCấp JWT 30 phút\nCông bố JWKS", "FFF4D6"),
        "api": (1145, 320, 1530, 590, "INTEGRATIONHUB API", "Lấy PUBLIC KEY/JWKS\nKiểm chữ ký, exp, aud\nKiểm scope finance.read", "EAF5EF"),
        "db": (1145, 700, 1530, 865, "SQL SERVER", "dbo.vw_Tin_Dung", "F4EAEA"),
    }
    for _, (x1, y1, x2, y2, title, detail, fill) in boxes.items():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill="#" + fill, outline="#2E74B5", width=3)
        tw = draw.textbbox((0, 0), title, font=bold)[2]
        draw.text(((x1 + x2 - tw) / 2, y1 + 32), title, font=bold, fill="#16324F")
        yy = y1 + 100
        for line in detail.split("\n"):
            draw.text((x1 + 38, yy), line, font=regular, fill="#202124")
            yy += 48

    def arrow(a, b, label, offset=(0, 0)):
        draw.line((a[0], a[1], b[0], b[1]), fill="#667085", width=5)
        dx, dy = b[0] - a[0], b[1] - a[1]
        if abs(dx) > abs(dy):
            tip = [(b[0], b[1]), (b[0] - 18 if dx > 0 else b[0] + 18, b[1] - 10), (b[0] - 18 if dx > 0 else b[0] + 18, b[1] + 10)]
        else:
            tip = [(b[0], b[1]), (b[0] - 10, b[1] - 18 if dy > 0 else b[1] + 18), (b[0] + 10, b[1] - 18 if dy > 0 else b[1] + 18)]
        draw.polygon(tip, fill="#667085")
        mx, my = (a[0] + b[0]) / 2 + offset[0], (a[1] + b[1]) / 2 + offset[1]
        draw.text((mx, my), label, font=small, fill="#1F4D78")

    arrow((470, 390), (605, 285), "client credentials", (-55, -42))
    arrow((605, 330), (470, 465), "JWT access_token", (-60, 18))
    arrow((470, 545), (1145, 545), "Authorization: Bearer ...", (110, 12))
    arrow((1145, 390), (1030, 290), "JWKS/public key", (-110, -48))
    arrow((1335, 590), (1335, 700), "SQL tham số hóa", (18, -5))
    draw.text((80, 25), "Private key never leaves the Authorization Server", font=bold, fill="#9B1C1C")
    img.save(path)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = MONO
        code.font.size = Pt(8.5)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(hp.add_run("INTEGRATIONHUB API  |  TÀI LIỆU KỸ THUẬT"), size=8.5, color=MID_GRAY, bold=True)
    footer = section.footer
    add_page_number(footer.paragraphs[0])
    # Built-in abstractNum 7 is Word's decimal List Number definition.
    # Separate num instances restart each logical sequence at 1.
    add_number_instance(doc, 10, 7)
    add_number_instance(doc, 11, 7)
    add_number_instance(doc, 12, 7)


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    set_run_font(p.add_run("HƯỚNG DẪN KIẾN TRÚC VÀ XÁC THỰC"), size=11, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run("API HẠN MỨC TÍN DỤNG"), size=28, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("IntegrationHub API · MSSQL dbo.vw_Tin_Dung"), size=15, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(40)
    set_run_font(p.add_run("Luồng xử lý từng class · JWT RSA private/public key · Token 30 phút · Hướng dẫn tích hợp"), size=10.5, color=MID_GRAY, italic=True)

    add_callout(doc, "TRẠNG THÁI HIỆN TẠI", "Endpoint lấy hạn mức tín dụng và kiểm tra scope finance.read đã có trong mã nguồn. Dịch vụ cấp token/Identity Provider chưa nằm trong repository; các địa chỉ token trong tài liệu phải được thay bằng cấu hình môi trường thực tế.", "info")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(64)
    set_run_font(p.add_run("Phiên bản 1.0  |  20/08/2026"), size=10, color=MID_GRAY)
    doc.add_page_break()


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    layer_img = ASSET_DIR / "architecture_layers.png"
    auth_img = ASSET_DIR / "authentication_flow.png"
    create_layer_diagram(layer_img)
    create_auth_diagram(auth_img)

    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "Mục lục nội dung", 1)
    for item in [
        "1. Phạm vi và cách đọc tài liệu",
        "2. Tổng quan kiến trúc API",
        "3. Trình tự xử lý một request từ class nào đến class nào",
        "4. Giải thích chi tiết từng model và class",
        "5. Hướng dẫn sử dụng API hạn mức tín dụng",
        "6. Kiến trúc xác thực JWT bằng private key/public key",
        "7. Quy trình lấy và tái sử dụng access token",
        "8. Quy tắc hết hạn 30 phút và cấu hình đề xuất",
        "9. Triển khai, vận hành và kiểm thử",
        "10. Phụ lục ví dụ mã nguồn",
    ]:
        p = add_para(doc, item)
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)

    add_heading(doc, "1. Phạm vi và cách đọc tài liệu", 1)
    add_para(doc, "Tài liệu này giải thích endpoint GET /api/v1/finance/credit-limits trong dự án api_core_system. Mục tiêu là giúp người phát triển hiểu cấu trúc nhiều lớp, luồng thực thi, dữ liệu vào/ra, cách truy vấn SQL Server và cách bảo vệ API bằng access token JWT.")
    add_callout(doc, "KẾT LUẬN NGẮN", "Client không gọi trực tiếp SQL Server. Client lấy access token từ Authorization Server, gửi token đến IntegrationHub API, API kiểm tra token và scope finance.read, sau đó mới gọi dbo.vw_Tin_Dung bằng câu lệnh SQL có tham số.", "ok")
    add_heading(doc, "1.1 Những gì đã có trong mã nguồn", 2)
    for text in [
        "Endpoint GET /api/v1/finance/credit-limits với hai query parameter ma_dt và ngay_hl.",
        "Policy authorization finance.read.",
        "JWT Bearer authentication dựa trên Authentication:Authority và Authentication:Audience.",
        "Application service kiểm tra dữ liệu đầu vào và chuẩn hóa ngày.",
        "Repository Dapper truy vấn dbo.vw_Tin_Dung bằng parameter, không ghép chuỗi SQL.",
        "Domain model CreditLimit và JSON camelCase.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "1.2 Những gì chưa có", 2)
    for text in [
        "Không có endpoint đăng nhập/cấp token trong IntegrationHub API.",
        "Không có private key trong repository và không nên thêm private key vào API này.",
        "Authority hiện là giá trị mẫu identity.example.com; cần cấu hình Identity Provider thật.",
        "Quy tắc token tối đa 30 phút chưa được cưỡng chế độc lập tại API; hiện API tin vào exp do Authority phát hành và cho ClockSkew 1 phút.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "2. Tổng quan kiến trúc API", 1)
    add_para(doc, "Dự án áp dụng kiến trúc phân lớp. Mỗi lớp chỉ giữ một loại trách nhiệm. Điều này giúp thay đổi database, quy tắc nghiệp vụ hoặc giao thức HTTP mà không làm toàn bộ mã nguồn phụ thuộc lẫn nhau.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(layer_img), width=Inches(6.25))
    picture._inline.docPr.set("descr", "Sơ đồ năm lớp API, Application, Domain, Infrastructure và Legacy SQL Server")
    picture._inline.docPr.set("title", "Kiến trúc phân lớp API hạn mức tín dụng")
    p.paragraph_format.space_after = Pt(3)
    p = doc.add_paragraph("Hình 1. Các lớp chính của use case hạn mức tín dụng")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs:
        set_run_font(r, size=9, color=MID_GRAY, italic=True)

    add_table(doc, ["Lớp", "Trách nhiệm", "Thành phần chính"], [
        ("API", "Nhận HTTP, binding tham số, yêu cầu quyền, trả HTTP response", "CreditLimitsController, middleware auth"),
        ("Application", "Điều phối use case, validate, định nghĩa cổng repository", "CreditLimitService, GetCreditLimitsQuery, ICreditLimitReadRepository"),
        ("Domain", "Mô hình dữ liệu nghiệp vụ độc lập", "CreditLimit"),
        ("Infrastructure", "Kết nối công nghệ bên ngoài và thực thi SQL", "CreditLimitReadRepository, SqlServerConnectionFactory"),
        ("Legacy SQL", "Nguồn dữ liệu thực", "dbo.vw_Tin_Dung"),
    ], [1500, 3760, 4100], 8.5)

    add_heading(doc, "2.1 Dependency Injection khi khởi động", 2)
    add_para(doc, "Program.cs gọi AddApplication() và AddInfrastructure(configuration). Hai extension method đăng ký các class vào DI container:")
    add_code(doc, "AddScoped<CreditLimitService>();\nAddScoped<ICreditLimitReadRepository, CreditLimitReadRepository>();\nAddSingleton<ILegacyDbConnectionFactory, SqlServerConnectionFactory>();")
    add_para(doc, "Khi ASP.NET Core cần tạo CreditLimitsController, DI tự tạo CreditLimitService; khi service cần ICreditLimitReadRepository, DI cấp CreditLimitReadRepository; repository nhận tiếp connection factory và database options. Đây là constructor injection.")

    add_heading(doc, "3. Trình tự xử lý một request", 1)
    add_callout(doc, "REQUEST MẪU", "GET /api/v1/finance/credit-limits?ma_dt=1000001&ngay_hl=2026-08-20 với header Authorization: Bearer <access_token>.", "info")
    steps = [
        ("Client gửi HTTPS request", "Query string mang ma_dt và ngay_hl; access token nằm trong header Authorization."),
        ("UseAuthentication() chạy", "JwtBearerHandler đọc Bearer token và tải metadata/JWKS từ Authority khi cần."),
        ("JWT được kiểm tra", "Chữ ký, issuer, audience và thời gian exp/nbf được kiểm tra. Token sai trả 401 trước khi controller chạy."),
        ("UseAuthorization() chạy", "Policy finance.read kiểm tra claim scope. Thiếu scope trả 403."),
        ("ASP.NET Core chọn controller", "Route api/v1/finance/credit-limits ánh xạ đến CreditLimitsController.Get()."),
        ("Model binding", "ma_dt được gán vào partnerCode; ngay_hl được parse thành DateTime? effectiveDate."),
        ("Controller tạo query model", "new GetCreditLimitsQuery(partnerCode, effectiveDate)."),
        ("Controller gọi application service", "CreditLimitService.GetAsync(query, cancellationToken)."),
        ("Service validate", "Bắt buộc ma_dt, tối đa 20 ký tự, bắt buộc ngay_hl; Trim mã và lấy Date của ngày hiệu lực."),
        ("Service gọi repository interface", "ICreditLimitReadRepository.GetAsync(...). Service không biết Dapper hay SQL Server."),
        ("Repository mở kết nối", "SqlServerConnectionFactory tạo SqlConnection từ LegacyDatabase:ConnectionString."),
        ("Repository chạy SQL", "Dapper truyền @PartnerCode và @EffectiveDate dưới dạng parameter có kiểu dữ liệu."),
        ("Dapper ánh xạ kết quả", "Alias cột SQL được map vào constructor của CreditLimit."),
        ("Controller trả response", "Ok(list) tạo HTTP 200; JSON serializer đổi tên property sang camelCase."),
    ]
    for title, detail in steps:
        add_step(doc, title, detail)
    add_code(doc, "Client\n  -> JWT middleware\n  -> finance.read policy\n  -> CreditLimitsController.Get\n  -> CreditLimitService.GetAsync\n  -> ICreditLimitReadRepository.GetAsync\n  -> CreditLimitReadRepository.GetAsync\n  -> SqlServerConnectionFactory.OpenConnectionAsync\n  -> dbo.vw_Tin_Dung\n  -> CreditLimit[] -> HTTP 200 JSON")

    add_heading(doc, "4. Giải thích chi tiết từng model và class", 1)
    add_heading(doc, "4.1 CreditLimitsController - lớp HTTP", 2)
    add_para(doc, "File: src/IntegrationHub.Api/Controllers/CreditLimitsController.cs")
    for text in [
        "[Route] định nghĩa URL của endpoint.",
        "[Authorize(Policy = ApiScopes.FinanceRead)] yêu cầu token có scope finance.read.",
        "[FromQuery(Name = \"ma_dt\")] và ngay_hl ánh xạ đúng tên parameter mà hệ thống tích hợp gửi.",
        "Controller chỉ chuyển đổi HTTP thành use case; không viết SQL trong controller.",
        "ArgumentException từ service được chuyển thành HTTP 400 ValidationProblemDetails.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "4.2 GetCreditLimitsQuery - input model", 2)
    add_code(doc, "public sealed record GetCreditLimitsQuery(\n    string? PartnerCode,\n    DateTime? EffectiveDate);")
    add_para(doc, "Đây là dữ liệu đầu vào nội bộ của use case. Dùng record giúp object bất biến theo ý nghĩa: controller tạo một giá trị query rồi truyền xuống service.")
    add_heading(doc, "4.3 CreditLimitService - quy tắc ứng dụng", 2)
    add_para(doc, "Service kiểm tra dữ liệu, tránh để controller hoặc repository gánh quy tắc nghiệp vụ. Nó trim PartnerCode, giới hạn 20 ký tự theo cột Ma_Dt varchar(20), yêu cầu EffectiveDate và chuẩn hóa Value.Date để bỏ phần giờ.")
    add_callout(doc, "LƯU Ý NGHIỆP VỤ", "Điều kiện hiện tại chỉ là Ngay_Kt >= ngay_hl. Nó chưa kiểm tra Ngay_Bd <= ngay_hl. Đây là đúng theo câu SQL yêu cầu ban đầu; nếu muốn 'còn hiệu lực tại một ngày', thường cần đồng thời Ngay_Bd <= ngay_hl.", "warn")
    add_heading(doc, "4.4 ICreditLimitReadRepository - abstraction", 2)
    add_para(doc, "Interface nằm ở Application, mô tả điều Application cần mà không chỉ ra công nghệ thực hiện. Nhờ đó có thể thay repository MSSQL bằng mock khi test hoặc nguồn dữ liệu khác mà service không đổi.")
    add_heading(doc, "4.5 CreditLimitReadRepository - SQL adapter", 2)
    add_para(doc, "Repository chứa câu SQL cố định, tạo DynamicParameters, mở connection và gọi Dapper QueryAsync<CreditLimit>. Parameterized SQL ngăn ma_dt bị hiểu như một đoạn SQL và giúp SQL Server tái sử dụng execution plan.")
    add_code(doc, "WHERE Ngay_Kt >= @EffectiveDate\n  AND Ma_Dt = @PartnerCode")
    add_heading(doc, "4.6 SqlServerConnectionFactory", 2)
    add_para(doc, "Factory đọc connection string từ options, tạo Microsoft.Data.SqlClient.SqlConnection và mở kết nối bất đồng bộ. await using bảo đảm connection được dispose sau khi query hoàn tất hoặc phát sinh lỗi.")

    add_heading(doc, "4.7 CreditLimit - output/domain model", 2)
    add_table(doc, ["Cột SQL", "Property C# / JSON", "Ý nghĩa"], [
        ("Ma_Dt", "PartnerCode / partnerCode", "Mã đối tượng/khách hàng"),
        ("Ma_Hd", "ContractCode / contractCode", "Mã hợp đồng"),
        ("Tien_Tin_Chap", "UnsecuredAmount / unsecuredAmount", "Hạn mức tín chấp"),
        ("Tien_Bao_Lanh", "GuaranteeAmount / guaranteeAmount", "Hạn mức bảo lãnh"),
        ("Tien_Cam_Co", "CollateralAmount / collateralAmount", "Hạn mức cầm cố"),
        ("Tien_Tin_Chap_Nt", "UnsecuredForeignCurrencyAmount", "Tín chấp nguyên tệ"),
        ("Tien_Bao_Lanh_Nt", "GuaranteeForeignCurrencyAmount", "Bảo lãnh nguyên tệ"),
        ("Tien_Cam_Co_Nt", "CollateralForeignCurrencyAmount", "Cầm cố nguyên tệ"),
        ("Ngay_Bd", "StartDate / startDate", "Ngày bắt đầu"),
        ("Ngay_Kt", "EndDate / endDate", "Ngày kết thúc"),
    ], [2100, 3600, 3660], 8.2)
    add_para(doc, "Các cột money được CONVERT(decimal(19,4)) trước khi Dapper map sang decimal để API có kiểu số rõ ràng. JSON serializer được cấu hình camelCase trong Program.cs.")

    add_heading(doc, "5. Hướng dẫn sử dụng API hạn mức tín dụng", 1)
    add_heading(doc, "5.1 Thông tin bắt buộc", 2)
    add_table(doc, ["Thông tin", "Giá trị", "Bắt buộc"], [
        ("Base URL", "Ví dụ https://api.company.vn", "Có - theo môi trường"),
        ("HTTP method", "GET", "Có"),
        ("Path", "/api/v1/finance/credit-limits", "Có"),
        ("ma_dt", "Mã đối tượng, tối đa 20 ký tự", "Có"),
        ("ngay_hl", "Ngày ISO yyyy-MM-dd", "Có"),
        ("Authorization", "Bearer <access_token>", "Có"),
        ("Scope trong token", "finance.read", "Có"),
    ], [2200, 5000, 2160], 8.7)
    add_heading(doc, "5.2 Ví dụ cURL", 2)
    add_code(doc, "curl --request GET \\\n  'https://api.company.vn/api/v1/finance/credit-limits?ma_dt=1000001&ngay_hl=2026-08-20' \\\n  --header 'Authorization: Bearer eyJhbGciOiJSUzI1NiIs...' \\\n  --header 'Accept: application/json'")
    add_heading(doc, "5.3 Ví dụ response 200", 2)
    add_code(doc, "[\n  {\n    \"partnerCode\": \"1000001\",\n    \"contractCode\": \"HD-001\",\n    \"unsecuredAmount\": 500000000.0000,\n    \"guaranteeAmount\": 0.0000,\n    \"collateralAmount\": 200000000.0000,\n    \"unsecuredForeignCurrencyAmount\": 0.0000,\n    \"guaranteeForeignCurrencyAmount\": 0.0000,\n    \"collateralForeignCurrencyAmount\": 0.0000,\n    \"startDate\": \"2026-01-01T00:00:00\",\n    \"endDate\": \"2026-12-31T00:00:00\"\n  }\n]")
    add_callout(doc, "KẾT QUẢ RỖNG", "Nếu không có bản ghi phù hợp, API trả HTTP 200 với mảng rỗng []. Đây không phải lỗi 404.", "info")

    add_heading(doc, "5.4 Mã trạng thái cần xử lý", 2)
    add_table(doc, ["HTTP", "Ý nghĩa", "Cách xử lý phía client"], [
        ("200", "Thành công; có thể là []", "Đọc JSON và xử lý danh sách"),
        ("400", "Thiếu/sai ma_dt hoặc ngay_hl", "Sửa request; không xin token lại"),
        ("401", "Không có token, token sai chữ ký hoặc hết hạn", "Xin token mới rồi retry một lần"),
        ("403", "Token hợp lệ nhưng thiếu finance.read", "Cấp đúng scope; xin token mới"),
        ("500", "Lỗi không xử lý trong API", "Ghi correlation ID/log, retry có kiểm soát"),
        ("503", "Phụ thuộc chưa sẵn sàng nếu được cấu hình", "Backoff; kiểm tra health/readiness"),
    ], [900, 3600, 4860], 8.5)
    add_heading(doc, "5.5 Ví dụ C# HttpClient", 2)
    add_code(doc, "using var request = new HttpRequestMessage(\n    HttpMethod.Get,\n    $\"{baseUrl}/api/v1/finance/credit-limits?ma_dt={Uri.EscapeDataString(maDt)}&ngay_hl={ngayHl:yyyy-MM-dd}\");\nrequest.Headers.Authorization = new AuthenticationHeaderValue(\"Bearer\", accessToken);\nusing var response = await httpClient.SendAsync(request, cancellationToken);\nresponse.EnsureSuccessStatusCode();")

    add_heading(doc, "6. Kiến trúc xác thực JWT bằng private/public key", 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(auth_img), width=Inches(6.35))
    picture._inline.docPr.set("descr", "Sơ đồ client xin JWT từ Authorization Server rồi gọi IntegrationHub API và SQL Server")
    picture._inline.docPr.set("title", "Luồng xác thực JWT private key public key")
    p.paragraph_format.space_after = Pt(3)
    p = doc.add_paragraph("Hình 2. Luồng cấp token, kiểm tra JWT và truy vấn dữ liệu")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        set_run_font(r, size=9, color=MID_GRAY, italic=True)

    add_heading(doc, "6.1 Vai trò của hai loại khóa", 2)
    add_table(doc, ["Khóa", "Nơi lưu", "Mục đích", "Có được chia sẻ?"], [
        ("Private key", "Authorization Server / HSM / Key Vault", "Ký JWT access token", "Không; tuyệt đối không đưa vào API/client/repository"),
        ("Public key", "Công bố qua JWKS; API cache", "Kiểm tra chữ ký JWT", "Có; đây không phải bí mật"),
    ], [1500, 2900, 2700, 2260], 8.5)
    add_para(doc, "Cơ chế này thường dùng RS256 hoặc PS256. Authorization Server ký phần header.payload bằng private key. IntegrationHub API dùng public key tương ứng với kid trong JWT để xác nhận token không bị sửa và đúng nơi phát hành.")
    add_callout(doc, "NGUYÊN TẮC", "IntegrationHub API không cần và không nên có private signing key nếu nó chỉ là resource server. Nếu private key của Authorization Server bị lộ, kẻ tấn công có thể tự tạo token hợp lệ.", "risk")

    add_heading(doc, "6.2 Claims tối thiểu nên có", 2)
    add_table(doc, ["Claim", "Ví dụ", "Mục đích"], [
        ("iss", "https://identity.company.vn/realms/integration", "Định danh nơi cấp token"),
        ("aud", "integration-api", "Token dành cho API nào"),
        ("sub/client_id", "odoo-integration", "Định danh client gọi"),
        ("scope", "finance.read", "Quyền được phép"),
        ("iat", "Unix timestamp", "Thời điểm phát hành"),
        ("nbf", "Unix timestamp", "Không hợp lệ trước thời điểm này"),
        ("exp", "iat + 1800 giây", "Hết hạn sau tối đa 30 phút"),
        ("jti", "UUID ngẫu nhiên", "Định danh duy nhất của token"),
    ], [1400, 3300, 4660], 8.5)

    add_heading(doc, "6.3 Kiểm tra trong mã nguồn hiện tại", 2)
    add_para(doc, "ScopeAuthorization.AddApiAuthentication() dùng AddJwtBearer. Authority cho middleware biết issuer và địa chỉ discovery; Audience là integration-api. Middleware tự tải OpenID Connect metadata và JWKS public keys, sau đó kiểm tra Bearer token.")
    add_code(doc, "options.Authority = authority;\noptions.Audience = audience;\noptions.RequireHttpsMetadata = true;\noptions.MapInboundClaims = false;\noptions.TokenValidationParameters.ClockSkew = TimeSpan.FromMinutes(1);")
    add_para(doc, "HasScope() đọc tất cả claim scope, tách theo khoảng trắng và so sánh chính xác với finance.read. Token có thể chứa nhiều scope, ví dụ \"master.read finance.read\".")

    add_heading(doc, "6.4 Luân chuyển khóa (key rotation)", 2)
    for text in [
        "Authorization Server tạo cặp khóa mới và thêm public key mới vào JWKS.",
        "Bắt đầu ký token mới bằng private key mới; JWT header mang kid mới.",
        "Giữ public key cũ trong JWKS ít nhất bằng thời gian sống token cộng độ lệch đồng hồ.",
        "Sau khi mọi token cũ hết hạn, mới loại public key cũ khỏi JWKS.",
        "Không đổi khóa đột ngột nếu không muốn toàn bộ token đang dùng bị 401.",
    ]:
        add_step(doc, "Bước", text, 11)

    add_heading(doc, "7. Quy trình lấy và tái sử dụng access token", 1)
    add_heading(doc, "7.1 Khuyến nghị: OAuth 2.0 Client Credentials", 2)
    add_para(doc, "Luồng này phù hợp cho Odoo/backend gọi backend, không có người dùng tương tác. Client gửi client_id và bí mật xác thực đến token endpoint của Authorization Server. Authorization Server trả access_token có scope finance.read.")
    add_callout(doc, "ĐỊA CHỈ TOKEN", "Repository hiện không định nghĩa token endpoint. Hãy đọc discovery document tại {Authority}/.well-known/openid-configuration và lấy trường token_endpoint. Nếu Authority thực tế là Keycloak, đường dẫn thường có dạng {Authority}/protocol/openid-connect/token.", "warn")
    add_code(doc, "curl --request POST '[TOKEN_ENDPOINT]' \\\n  --header 'Content-Type: application/x-www-form-urlencoded' \\\n  --data-urlencode 'grant_type=client_credentials' \\\n  --data-urlencode 'client_id=odoo-integration' \\\n  --data-urlencode 'client_secret=[CLIENT_SECRET]' \\\n  --data-urlencode 'scope=finance.read'")
    add_heading(doc, "7.2 Response token điển hình", 2)
    add_code(doc, "{\n  \"access_token\": \"eyJhbGciOiJSUzI1NiIsImtpZCI6Ii4uLiJ9...\",\n  \"token_type\": \"Bearer\",\n  \"expires_in\": 1800,\n  \"scope\": \"finance.read\"\n}")
    add_heading(doc, "7.3 Cách lưu và tái sử dụng", 2)
    for title, detail in [
        ("Lưu trong bộ nhớ", "Backend giữ access_token và expires_at = thời điểm nhận + expires_in. Không ghi token vào log."),
        ("Tái sử dụng", "Dùng cùng token cho nhiều request trong thời gian hợp lệ; không cần xin token cho từng request."),
        ("Làm mới sớm", "Khi còn dưới 60 giây, xin token mới trước request tiếp theo để tránh hết hạn giữa đường."),
        ("Gặp 401", "Xóa token cache, xin token mới và retry đúng một lần; không retry vô hạn."),
        ("Đồng bộ thời gian", "Client, Authorization Server và API phải dùng NTP; sai giờ gây token chưa có hiệu lực hoặc hết hạn sớm."),
    ]:
        add_step(doc, title, detail, 12)

    add_heading(doc, "7.4 Client authentication bằng private_key_jwt (nâng cao)", 2)
    add_para(doc, "Nếu không muốn dùng client_secret, mỗi client có thể giữ private key riêng để ký client_assertion khi xin token. Authorization Server lưu public key của client để xác minh. Đây là cặp khóa của client, khác với cặp khóa Authorization Server dùng để ký access token.")
    add_code(doc, "POST [TOKEN_ENDPOINT]\nContent-Type: application/x-www-form-urlencoded\n\ngrant_type=client_credentials\nclient_id=odoo-integration\nscope=finance.read\nclient_assertion_type=urn:ietf:params:oauth:client-assertion-type:jwt-bearer\nclient_assertion=[JWT_DO_CLIENT_KY_BANG_PRIVATE_KEY_CUA_CLIENT]")
    add_table(doc, ["JWT client_assertion", "Giá trị yêu cầu"], [
        ("iss", "client_id"),
        ("sub", "client_id"),
        ("aud", "token endpoint chính xác"),
        ("exp", "Ngắn, thường không quá vài phút"),
        ("jti", "Duy nhất để chống replay"),
    ], [2600, 6760], 9)
    add_callout(doc, "PHÂN BIỆT", "Private key của client chỉ dùng chứng minh danh tính client khi xin token. Private key của Authorization Server dùng ký access token. IntegrationHub API chỉ cần public key/JWKS của Authorization Server.", "info")

    add_heading(doc, "8. Quy tắc hết hạn 30 phút", 1)
    add_heading(doc, "8.1 Cấu hình bắt buộc tại Authorization Server", 2)
    for text in [
        "Access token lifetime = 1800 giây.",
        "Token phải có iat và exp; exp - iat không lớn hơn 1800 giây.",
        "Không cấp scope finance.read cho client không được phép.",
        "Client Credentials thường không cần refresh token; hết hạn thì gọi token endpoint để lấy token mới.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "8.2 Cấu hình API để hết hạn đúng thời điểm", 2)
    add_para(doc, "Mã hiện tại đặt ClockSkew = 1 phút, nghĩa là middleware có thể chấp nhận token thêm tối đa khoảng một phút quanh mốc thời gian để bù sai lệch đồng hồ. Nếu yêu cầu nghiệp vụ là hết hạn chính xác sau 30 phút, đổi về TimeSpan.Zero và bắt buộc đồng bộ NTP.")
    add_code(doc, "options.TokenValidationParameters.ClockSkew = TimeSpan.Zero;")
    add_callout(doc, "QUAN TRỌNG", "Chỉ kiểm tra exp chưa ngăn Authorization Server vô tình cấp token sống 2 giờ. Để bảo vệ kép, API có thể thêm LifetimeValidator kiểm tra exp - iat <= 30 phút. Tuy nhiên cấu hình chuẩn nhất vẫn là đặt lifetime tại Authorization Server.", "warn")
    add_heading(doc, "8.3 Ma trận thời gian", 2)
    add_table(doc, ["Thời điểm", "Trạng thái", "Hành động"], [
        ("T0", "Client nhận token expires_in=1800", "Cache token và expires_at"),
        ("T0 + 0..28:59", "Token dùng bình thường", "Tái sử dụng"),
        ("T0 + 29:00", "Vùng làm mới sớm 60 giây", "Xin token mới trước request tiếp theo"),
        ("T0 + 30:00", "Token hết hạn", "API trả 401 nếu ClockSkew=0"),
        ("Sau 401", "Token cache không còn dùng được", "Xin mới và retry một lần"),
    ], [2100, 3900, 3360], 8.5)

    add_heading(doc, "9. Triển khai, vận hành và kiểm thử", 1)
    add_heading(doc, "9.1 Cấu hình API", 2)
    add_code(doc, "Authentication__Authority=https://identity.company.vn/realms/integration\nAuthentication__Audience=integration-api\nLegacyDatabase__ConnectionString=Server=...;Database=...;User ID=integration_api;Password=...;Encrypt=True;TrustServerCertificate=False")
    add_para(doc, "Không commit connection string, client secret, private key hoặc access token. Dùng secret store của môi trường triển khai. API bắt buộc HTTPS và Authority cũng phải dùng HTTPS.")
    add_heading(doc, "9.2 Quyền database tối thiểu", 2)
    add_code(doc, "GRANT SELECT ON OBJECT::dbo.vw_Tin_Dung TO integration_api;")
    add_para(doc, "Tài khoản integration_api chỉ cần SELECT view. Không cấp db_owner hoặc quyền ghi nếu use case chỉ đọc.")
    add_heading(doc, "9.3 Bộ test chấp nhận", 2)
    tests = [
        ("Không gửi token", "401"),
        ("Token chữ ký sai", "401"),
        ("Token đúng nhưng aud khác", "401"),
        ("Token hết hạn 30 phút", "401"),
        ("Token thiếu finance.read", "403"),
        ("Thiếu ma_dt hoặc ngay_hl", "400"),
        ("ngay_hl sai định dạng", "400"),
        ("ma_dt dài hơn 20 ký tự", "400"),
        ("Không có dữ liệu", "200 và []"),
        ("Có dữ liệu", "200 và JSON đúng mapping"),
        ("ma_dt chứa ký tự SQL injection", "Không làm thay đổi câu SQL; chỉ là giá trị parameter"),
    ]
    add_table(doc, ["Kịch bản", "Kết quả mong đợi"], tests, [6000, 3360], 8.8)
    add_heading(doc, "9.4 Logging an toàn", 2)
    for text in [
        "Ghi request ID/correlation ID, client_id/sub, route, status code và thời gian xử lý.",
        "Không ghi Authorization header, access token, private key, client secret hoặc connection string.",
        "Có thể ghi ma_dt theo chính sách dữ liệu nội bộ; cân nhắc masking nếu đây là thông tin nhạy cảm.",
        "Theo dõi riêng số lượng 401, 403, lỗi kết nối MSSQL và query timeout.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "10. Phụ lục ví dụ mã nguồn", 1)
    add_heading(doc, "10.1 Pseudocode token cache phía client", 2)
    add_code(doc, "if (cachedToken is null || UtcNow >= cachedToken.ExpiresAt - 60 seconds)\n{\n    cachedToken = await RequestClientCredentialsToken();\n}\n\nresponse = await CallCreditLimitApi(cachedToken.AccessToken);\n\nif (response.StatusCode == 401 && notRetried)\n{\n    cachedToken = await RequestClientCredentialsToken();\n    response = await CallCreditLimitApi(cachedToken.AccessToken);\n}")
    add_heading(doc, "10.2 Checklist trước khi go-live", 2)
    for text in [
        "Authority và Audience đã trỏ đúng môi trường.",
        "Authorization Server ký JWT bằng RSA; private key ở HSM/Key Vault hoặc secret store bảo vệ.",
        "JWKS truy cập được từ API và có quy trình rotation.",
        "Access token lifetime = 1800 giây; ClockSkew được chọn phù hợp yêu cầu.",
        "Client được cấp đúng finance.read và không có scope thừa.",
        "Database user chỉ có SELECT dbo.vw_Tin_Dung.",
        "HTTPS, NTP, logging redaction và monitoring đã bật.",
        "Đã chạy toàn bộ test 401/403/400/200 và hết hạn 30 phút.",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "ĐỀ XUẤT BƯỚC TIẾP THEO", "Chọn Authorization Server cụ thể (ví dụ Keycloak hoặc dịch vụ OIDC nội bộ), cấu hình realm/client/scope và token lifetime 1800 giây; sau đó cập nhật Authority thật, đặt ClockSkew theo yêu cầu và bổ sung integration test dùng token đã ký bằng key thử nghiệm.", "ok")

    add_heading(doc, "Tài liệu mã nguồn tham chiếu", 1)
    for text in [
        "src/IntegrationHub.Api/Program.cs",
        "src/IntegrationHub.Api/Authentication/ScopeAuthorization.cs",
        "src/IntegrationHub.Api/Authentication/ApiScopes.cs",
        "src/IntegrationHub.Api/Controllers/CreditLimitsController.cs",
        "src/IntegrationHub.Application/Modules/Finance/CreditLimits/CreditLimitService.cs",
        "src/IntegrationHub.Application/Modules/Finance/CreditLimits/GetCreditLimitsQuery.cs",
        "src/IntegrationHub.Application/Abstractions/ICreditLimitReadRepository.cs",
        "src/IntegrationHub.Infrastructure/LegacySqlServer/Finance/CreditLimitReadRepository.cs",
        "src/IntegrationHub.Infrastructure/LegacySqlServer/SqlServerConnectionFactory.cs",
        "src/IntegrationHub.Domain/Finance/CreditLimit.cs",
    ]:
        add_bullet(doc, text)

    core = doc.core_properties
    core.title = "Hướng dẫn kiến trúc và xác thực API hạn mức tín dụng"
    core.subject = "IntegrationHub API, MSSQL, JWT RSA và access token 30 phút"
    core.author = "Đội phát triển IntegrationHub"
    core.keywords = "IntegrationHub, API, JWT, RSA, private key, public key, MSSQL, vw_Tin_Dung"
    core.comments = "Tài liệu kỹ thuật nội bộ"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
